#!/usr/bin/env python3
"""
Prospection de commerces locaux sans site web (OpenStreetMap / Overpass).

Usage :
    python prospection.py --ville "Lyon" --rayon 5
    python prospection.py --ville "Lyon" --rayon 5 --no-agent      # sans enrichissement IA
    python prospection.py --ville "Lyon" --rayon 5 --max-agents 50 # limite le coût IA

Dépendances : pip install requests openpyxl anthropic
Pour l'agent IA : export ANTHROPIC_API_KEY=...
"""
import argparse
import json
import os
import re
import sys
import time
from concurrent.futures import ThreadPoolExecutor, as_completed

import requests
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

UA = {"User-Agent": "prospection-locale/1.0 (script personnel)"}
OVERPASS_URL = "https://overpass-api.de/api/interpreter"
NOMINATIM_URL = "https://nominatim.openstreetmap.org/search"
AGENT_MODEL = "claude-haiku-5-5"  # "claude-sonnet-5-5" si tu veux plus de finesse

# ÉTAPE 1 — Catégories : libellé -> filtres de tags OSM (syntaxe Overpass).
CATEGORIES = {
    "Restaurant": ['["amenity"="restaurant"]'],
    "Coiffeur": ['["shop"="hairdresser"]'],
    "Boulangerie": ['["shop"="bakery"]'],
    "Artisan": ['["craft"]'],
    "Garage": ['["shop"="car_repair"]', '["amenity"="car_repair"]'],
    "Institut de beauté": ['["shop"="beauty"]'],
}


# ÉTAPE 2 — Géocodage : transforme "Lyon" en coordonnées GPS (centre-ville).
def geocoder(ville):
    r = requests.get(NOMINATIM_URL, params={"q": ville, "format": "json", "limit": 1},
                     headers=UA, timeout=30)
    r.raise_for_status()
    data = r.json()
    if not data:
        sys.exit(f"Ville introuvable : {ville}")
    return float(data[0]["lat"]), float(data[0]["lon"])


# ÉTAPE 3 — Requête Overpass : on demande tous les éléments portant les tags voulus
# dans un cercle (around) de `rayon_km` autour du centre. "nwr" = nodes, ways, relations
# et "out center" donne un point pour les surfaces (bâtiments).
def requete_overpass(lat, lon, rayon_km):
    rayon_m = int(rayon_km * 1000)
    blocs = []
    for filtres in CATEGORIES.values():
        for f in filtres:
            blocs.append(f'nwr{f}["name"](around:{rayon_m},{lat},{lon});')
    query = f"[out:json][timeout:120];({''.join(blocs)});out center tags;"
    for tentative in range(3):  # Overpass est parfois surchargé : on réessaie
        r = requests.post(OVERPASS_URL, data={"data": query}, headers=UA, timeout=180)
        if r.status_code == 200:
            return r.json()["elements"]
        time.sleep(5 * (tentative + 1))
    r.raise_for_status()


# ÉTAPE 4 — Normalisation : on transforme un élément OSM brut en dict propre.
def categorie_de(tags):
    if tags.get("amenity") == "restaurant":
        return "Restaurant"
    shop = tags.get("shop")
    if shop == "hairdresser":
        return "Coiffeur"
    if shop == "bakery":
        return "Boulangerie"
    if shop == "beauty":
        return "Institut de beauté"
    if shop == "car_repair" or tags.get("amenity") == "car_repair":
        return "Garage"
    if "craft" in tags:
        return "Artisan"
    return None


def adresse_de(tags):
    rue = " ".join(x for x in (tags.get("addr:housenumber"), tags.get("addr:street")) if x)
    ville = " ".join(x for x in (tags.get("addr:postcode"), tags.get("addr:city")) if x)
    return ", ".join(x for x in (rue, ville) if x)


def quartier_de(tags):
    # Pas toujours renseigné dans OSM : on retombe sur le code postal.
    return (tags.get("addr:suburb") or tags.get("addr:neighbourhood")
            or tags.get("addr:postcode") or "")


def normaliser(elements):
    vus, out = set(), []
    for el in elements:
        t = el.get("tags", {})
        cat = categorie_de(t)
        if not cat:
            continue
        cle = (el["type"], el["id"])
        if cle in vus:
            continue
        vus.add(cle)
        site = t.get("website") or t.get("contact:website") or t.get("url") or ""
        tel = t.get("phone") or t.get("contact:phone") or t.get("contact:mobile") or ""
        out.append({
            "nom": t["name"], "categorie": cat, "adresse": adresse_de(t),
            "quartier": quartier_de(t), "telephone": tel, "site": site.strip(),
            "statut": "", "notes": "",
        })
    return out


# ÉTAPE 5 — Vérification du site : une requête HTTP ; si erreur réseau ou code >= 400
# (hors 401/403/429 : le site existe mais bloque les robots), on dit "site mort".
def verifier_site(url):
    if not url:
        return "Pas de site"
    if not re.match(r"^https?://", url):
        url = "http://" + url
    try:
        r = requests.get(url, headers=UA, timeout=10, allow_redirects=True, stream=True)
        r.close()
        if r.status_code < 400 or r.status_code in (401, 403, 429):
            return "OK"
        return "Site mort"
    except requests.RequestException:
        return "Site mort"


# ÉTAPE 6 — Agent IA (Haiku) : pour CHAQUE commerce, un mini-agent cherche sur le web
# (outil web_search d'Anthropic) le téléphone, le site officiel et une note utile
# pour la prospection. Il répond en JSON strict, qu'on fusionne dans nos données.
SYSTEM_AGENT = (
    "Tu es un agent de recherche pour un freelance qui prospecte des commerces locaux. "
    "Cherche sur le web des informations FIABLES sur le commerce donné. "
    "Réponds UNIQUEMENT par un objet JSON : "
    '{"telephone": "" , "site": "", "notes": ""}. '
    "'site' = site web officiel du commerce (pas Facebook, TripAdvisor, PagesJaunes, "
    "annuaires ; chaîne vide si aucun). 'notes' = 1 phrase max utile au démarchage "
    "(présence sur réseaux sociaux, avis, signes d'activité ou de fermeture). "
    "N'invente rien : laisse vide si tu n'es pas sûr."
)


def agent_enrichir(client, commerce, ville):
    prompt = (f"Commerce : {commerce['nom']}\nCatégorie : {commerce['categorie']}\n"
              f"Adresse : {commerce['adresse'] or 'inconnue'}\nVille : {ville}")
    resp = client.messages.create(
        model=AGENT_MODEL,
        max_tokens=600,
        system=SYSTEM_AGENT,
        tools=[{"type": "web_search_20250305", "name": "web_search", "max_uses": 2}],
        messages=[{"role": "user", "content": prompt}],
    )
    texte = "".join(b.text for b in resp.content if b.type == "text")
    m = re.search(r"\{.*\}", texte, re.S)
    return json.loads(m.group(0)) if m else {}


def enrichir_tous(commerces, ville, max_agents, workers=5):
    try:
        import anthropic
    except ImportError:
        print("Module 'anthropic' absent : enrichissement IA ignoré.")
        return
    if not os.environ.get("ANTHROPIC_API_KEY"):
        print("ANTHROPIC_API_KEY absente : enrichissement IA ignoré.")
        return
    client = anthropic.Anthropic()
    # On traite en priorité ceux qui manquent d'infos (pas de site / pas de téléphone).
    file = sorted(commerces, key=lambda c: (bool(c["site"]), bool(c["telephone"])))
    if max_agents:
        file = file[:max_agents]
    print(f"Agents {AGENT_MODEL} sur {len(file)} commerces…")

    def tache(c):
        try:
            return c, agent_enrichir(client, c, ville)
        except Exception as e:  # une erreur sur un commerce ne doit pas tout arrêter
            return c, {"notes": f"[agent en échec : {type(e).__name__}]"}

    with ThreadPoolExecutor(max_workers=workers) as ex:
        for i, f in enumerate(as_completed([ex.submit(tache, c) for c in file]), 1):
            c, info = f.result()
            # On ne remplace jamais une donnée OSM existante : l'agent complète seulement.
            if not c["telephone"] and info.get("telephone"):
                c["telephone"] = str(info["telephone"])
            if not c["site"] and info.get("site"):
                c["site"] = str(info["site"])
                c["site_trouve_par_agent"] = True
            if info.get("notes"):
                c["notes"] = str(info["notes"])
            if i % 10 == 0:
                print(f"  {i}/{len(file)}")


# ÉTAPE 7 — Export Excel : colonnes demandées, tri par priorité puis quartier,
# lignes "Haute" colorées, menu déroulant oui/non pour "Démarché".
COLONNES = ["Nom", "Catégorie", "Adresse", "Téléphone", "Site", "Statut site",
            "Priorité", "Notes", "Démarché (oui/non)"]


def exporter(commerces, chemin):
    ordre = {"Haute": 0, "Basse": 1}
    commerces.sort(key=lambda c: (ordre[c["priorite"]], c["quartier"].lower(), c["nom"].lower()))

    wb = Workbook()
    ws = wb.active
    ws.title = "Prospects"
    ws.append(COLONNES)
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="305496")

    rouge = PatternFill("solid", fgColor="FCE4D6")
    for c in commerces:
        note = c["notes"]
        if c.get("site_trouve_par_agent"):
            note = ("Site trouvé par l'agent (hors OSM). " + note).strip()
        ws.append([c["nom"], c["categorie"], c["adresse"], c["telephone"], c["site"],
                   c["statut"], c["priorite"], note, "non"])
        if c["priorite"] == "Haute":
            for cell in ws[ws.max_row]:
                cell.fill = rouge

    dv = DataValidation(type="list", formula1='"oui,non"', allow_blank=False)
    ws.add_data_validation(dv)
    dv.add(f"I2:I{ws.max_row}")

    largeurs = [32, 18, 40, 18, 35, 14, 10, 50, 18]
    for i, w in enumerate(largeurs, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    wb.save(chemin)


def main():
    p = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    p.add_argument("--ville", required=True, help="Ville centre de la zone")
    p.add_argument("--rayon", type=float, default=5, help="Rayon en km (défaut 5)")
    p.add_argument("--sortie", default="prospects.xlsx")
    p.add_argument("--no-agent", action="store_true", help="Désactive les agents IA")
    p.add_argument("--max-agents", type=int, default=0, help="Nb max de commerces enrichis (0 = tous)")
    a = p.parse_args()

    print(f"Géocodage de {a.ville}…")
    lat, lon = geocoder(a.ville)
    print(f"Centre : {lat:.4f}, {lon:.4f} — requête Overpass ({a.rayon} km)…")
    commerces = normaliser(requete_overpass(lat, lon, a.rayon))
    print(f"{len(commerces)} commerces trouvés.")

    if not a.no_agent:
        enrichir_tous(commerces, a.ville, a.max_agents)

    print("Vérification des sites…")
    urls = {c["site"] for c in commerces if c["site"]}
    with ThreadPoolExecutor(max_workers=16) as ex:
        resultats = dict(zip(urls, ex.map(verifier_site, urls)))
    for c in commerces:
        c["statut"] = resultats.get(c["site"], "Pas de site") if c["site"] else "Pas de site"
        c["priorite"] = "Haute" if c["statut"] in ("Pas de site", "Site mort") else "Basse"

    exporter(commerces, a.sortie)
    hautes = sum(c["priorite"] == "Haute" for c in commerces)
    print(f"Terminé : {a.sortie} ({hautes} prospects prioritaires sur {len(commerces)}).")


if __name__ == "__main__":
    main()
